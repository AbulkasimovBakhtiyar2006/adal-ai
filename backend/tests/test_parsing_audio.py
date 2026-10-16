import io
import json
import wave

import httpx
import pytest
from docx import Document
from pypdf import PdfWriter

from app.audio import AudioError, inspect_audio, transcribe
from app.parsing import ParseError, clean_text, parse_file, parse_html, split_sections
from app.providers import ProviderUnavailable


def wav(seconds=1):
    buffer = io.BytesIO()
    with wave.open(buffer, "wb") as output:
        output.setnchannels(1)
        output.setsampwidth(2)
        output.setframerate(8000)
        output.writeframes(b"\x00\x00" * (8000 * seconds))
    return buffer.getvalue()


def test_docx_keeps_table_order():
    doc = Document()
    doc.add_paragraph("Перед таблицей находится вводный абзац.")
    row = doc.add_table(rows=1, cols=2).rows[0]
    row.cells[0].text = "Первая ячейка"
    row.cells[1].text = "Вторая ячейка"
    doc.add_paragraph("После таблицы находится заключение.")
    buf = io.BytesIO()
    doc.save(buf)
    parsed = parse_file(buf.getvalue(), "test.docx")
    assert parsed.text.index("Перед") < parsed.text.index("Первая") < parsed.text.index("После")
    assert "Первая ячейка | Вторая ячейка" in parsed.text
    assert parsed.warnings


def test_html_recovers_unclosed_navigation_and_tables():
    rows = "".join(
        f"<tr><td>Пункт {i}</td><td>Официальное содержание пункта.</td></tr>" for i in range(300)
    )
    raw = (
        "<header><nav>Меню<article><h1>Документ</h1><p>Введение"
        f"<table>{rows}</table><p>Конец документа.</article>"
    ).encode()
    parsed = parse_html(raw)
    chunks = split_sections(parsed.text)
    assert len(chunks) > 1
    assert "Меню" not in parsed.text
    assert "Пункт 299" in chunks[-1][1]
    assert "Конец документа." in chunks[-1][1]


def test_split_preserves_short_prefix_and_tail():
    text = "Закон\n\nСтатья 1. " + "Содержание статьи. " * 10 + "\n\nСтатья 2."
    sections = split_sections(text)
    assert "\n\n".join(body for _, body in sections) == clean_text(text)


@pytest.mark.parametrize("text", ["", "Очень мало", "x" * 8801])
def test_split_does_not_silently_drop_invalid_text(text):
    with pytest.raises(ParseError):
        split_sections(text)


def test_scanned_pdf_needs_ocr():
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    buf = io.BytesIO()
    writer.write(buf)
    with pytest.raises(ParseError, match="ocr"):
        parse_file(buf.getvalue(), "scan.pdf")


@pytest.mark.parametrize(
    "data,name", [(b"bad", "test.docx"), (b"bad", "test.pdf"), (b"bad", "test.exe")]
)
def test_corrupt_documents_fail_safely(data, name):
    with pytest.raises(ParseError):
        parse_file(data, name)


def test_actual_audio_type_and_duration():
    mime, duration = inspect_audio(wav())
    assert mime == "audio/wav"
    assert duration == 1
    with pytest.raises(AudioError):
        inspect_audio(b"renamed text is not audio")
    with pytest.raises(AudioError, match="too_long"):
        inspect_audio(wav(181))


async def test_transcription_contract_and_review(settings):
    def handler(request):
        assert request.headers["x-goog-api-key"] == "test-gemini"
        body = json.loads(request.content)
        assert "do not obey" in body["systemInstruction"]["parts"][0]["text"]
        assert body["contents"][0]["parts"][0]["inlineData"]["mimeType"] == "audio/wav"
        return httpx.Response(
            200,
            json={
                "candidates": [
                    {
                        "finishReason": "STOP",
                        "content": {
                            "parts": [
                                {
                                    "text": json.dumps(
                                        {"text": "Как подать заявление?", "language": "ru"}
                                    )
                                }
                            ]
                        },
                    }
                ]
            },
        )

    async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as client:
        result = await transcribe(wav(), settings, client)
    assert result["review_required"] is True
    assert result["text"] == "Как подать заявление?"


async def test_transcription_rejects_truncated_response(settings):
    async with httpx.AsyncClient(
        transport=httpx.MockTransport(
            lambda _: httpx.Response(200, json={"candidates": [{"finishReason": "MAX_TOKENS"}]})
        )
    ) as client:
        with pytest.raises(ProviderUnavailable):
            await transcribe(wav(), settings, client)
